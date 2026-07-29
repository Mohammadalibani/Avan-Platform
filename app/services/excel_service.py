# app/services/excel_service.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO
from datetime import datetime
import jdatetime

class ExcelService:
    def __init__(self, template=None):
        self.template = template or {}
        self.set_default_template()
    
    def set_default_template(self):
        """تنظیم قالب پیش‌فرض"""
        self.template = {
            'header_bg_color': '2c3e50',
            'header_text_color': 'ffffff',
            'even_row_color': 'f8f9fa',
            'odd_row_color': 'ffffff',
            'outer_border_style': 'thick',
            'vertical_border_style': 'thin',
            'horizontal_border_style': 'dotted',
            'border_color': '000000',
            'font_name': 'B Nazanin',
            'header_font_size': 12,
            'data_font_size': 11,
        }
    
    def create_personnel_report(self, personnel_data, columns, title="گزارش پرسنل"):
        """ایجاد گزارش اکسل پرسنل"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title
        
        # تنظیمات فونت
        header_font = Font(
            name=self.template['font_name'],
            size=self.template['header_font_size'],
            bold=True,
            color=self.template['header_text_color']
        )
        data_font = Font(
            name=self.template['font_name'],
            size=self.template['data_font_size']
        )
        
        # تنظیمات رنگ‌ها
        header_fill = PatternFill(
            start_color=self.template['header_bg_color'],
            end_color=self.template['header_bg_color'],
            fill_type='solid'
        )
        even_fill = PatternFill(
            start_color=self.template['even_row_color'],
            end_color=self.template['even_row_color'],
            fill_type='solid'
        )
        odd_fill = PatternFill(
            start_color=self.template['odd_row_color'],
            end_color=self.template['odd_row_color'],
            fill_type='solid'
        )
        
        # تنظیمات خطوط
        border_style = self.template['outer_border_style']
        border_color = self.template['border_color']
        thin_border = Border(
            left=Side(style=border_style, color=border_color),
            right=Side(style=border_style, color=border_color),
            top=Side(style=border_style, color=border_color),
            bottom=Side(style=border_style, color=border_color)
        )
        
        # عنوان گزارش
        ws.merge_cells(f'A1:{get_column_letter(len(columns))}1')
        title_cell = ws['A1']
        title_cell.value = title
        title_cell.font = Font(
            name=self.template['font_name'],
            size=14,
            bold=True
        )
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # هدر ستون‌ها
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=2, column=col_idx)
            cell.value = col['title']
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = 20
        
        # داده‌ها
        for row_idx, record in enumerate(personnel_data, start=3):
            for col_idx, col in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = record.get(col['key'], '')
                
                # تبدیل تاریخ شمسی
                if col.get('type') == 'date' and value:
                    try:
                        if isinstance(value, str):
                            dt = datetime.fromisoformat(value)
                            value = jdatetime.datetime.fromgregorian(datetime=dt).strftime('%Y/%m/%d')
                    except:
                        pass
                
                cell.value = value
                cell.font = data_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
                
                # رنگ‌بندی ردیف‌ها
                if row_idx % 2 == 0:
                    cell.fill = even_fill
                else:
                    cell.fill = odd_fill
        
        # تنظیم ارتفاع ردیف‌ها
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            ws.row_dimensions[row[0].row].height = 25
        
        # ذخیره در BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
    
    def create_requests_report(self, requests_data, title="گزارش درخواست‌ها"):
        """ایجاد گزارش درخواست‌ها"""
        columns = [
            {'key': 'id', 'title': 'ردیف'},
            {'key': 'requester_name', 'title': 'درخواست‌دهنده'},
            {'key': 'request_type_persian', 'title': 'نوع درخواست'},
            {'key': 'status_persian', 'title': 'وضعیت'},
            {'key': 'request_date', 'title': 'تاریخ ثبت', 'type': 'date'},
            {'key': 'extra_data', 'title': 'اطلاعات تکمیلی'},
        ]
        return self.create_personnel_report(requests_data, columns, title)